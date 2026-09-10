#!/usr/bin/env python3
"""Refine the supplied PrismQuant layout as native editable PowerPoint objects.

Usage:
  python refine_prismquant.py --out .
  python refine_prismquant.py --out . --render --soffice /path/to/soffice
Dependencies: python-pptx, Pillow, PyMuPDF. Rendering: LibreOffice and Times New Roman.
The scatter and bar microdiagrams are conceptual, not experimental data.
"""
from __future__ import annotations
import argparse, json, math, os, re, shutil, subprocess
from functools import lru_cache
from PIL import ImageFont
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.oxml.xmlchemy import OxmlElement

C = dict(ink='233E5D', purple='644A86', blue='477F9F', had='A8DADC',
         weight='F4F7ED', panel='F3F5F8', cache='F7E7DC', cacheedge='C38F70',
         border='BBC6D2', lightpurple='EFEAF5', lightblue='E7F1F6',
         green='789989', softgreen='BED0C2', white='FFFFFF', muted='557087')
SCALE = .72  # logical 100 pixels per inch
FONT = 'Times New Roman'
prs = Presentation(); prs.slide_width = Inches(16); prs.slide_height = Inches(8)
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid(); slide.background.fill.fore_color.rgb = RGBColor.from_string(C['white'])
shapes = slide.shapes
panels = []
text_records = []

def u(v): return Inches(v / 100)
def rgb(v): return RGBColor.from_string(C.get(v, v))
def group(name):
    global shapes
    g=slide.shapes.add_group_shape(); g.name=name; shapes=g.shapes
    return g

def rect(x,y,w,h,fill='white',line='border',lw=.85,r=7,name=''):
    shape=shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if r else MSO_SHAPE.RECTANGLE,u(x),u(y),u(w),u(h))
    if r: shape.adjustments[0]=min(r/min(w,h),.45)
    if fill: shape.fill.solid(); shape.fill.fore_color.rgb=rgb(fill)
    else: shape.fill.background()
    if line: shape.line.color.rgb=rgb(line); shape.line.width=Pt(lw)
    else: shape.line.fill.background()
    if name: shape.name=name
    return shape

def ellipse(x,y,w,h,fill='white',line='ink',lw=1.05):
    s=shapes.add_shape(MSO_SHAPE.OVAL,u(x),u(y),u(w),u(h))
    if fill: s.fill.solid(); s.fill.fore_color.rgb=rgb(fill)
    else:s.fill.background()
    if line:s.line.color.rgb=rgb(line);s.line.width=Pt(lw)
    else:s.line.fill.background()
    return s

def line(points,color='ink',lw=1.25,arrow=True,dash=False):
    out=[]
    for i,(a,b) in enumerate(zip(points,points[1:])):
        s=shapes.add_connector(MSO_CONNECTOR.STRAIGHT,u(a[0]),u(a[1]),u(b[0]),u(b[1]))
        s.line.color.rgb=rgb(color);s.line.width=Pt(lw)
        ln=s._element.spPr.get_or_add_ln()
        if dash:
            el=OxmlElement('a:prstDash');el.set('val','dash');ln.append(el)
        if arrow and i==len(points)-2:
            el=OxmlElement('a:tailEnd');el.set('type','triangle');el.set('w','sm');el.set('len','sm');ln.append(el)
        out.append(s)
    return out

def dot(x,y,r=2.2,color='ink'):return ellipse(x-r,y-r,2*r,2*r,color,None)

def runs_math(s):
    """Editable rich runs for true subscripts/superscripts, with no TeX in PPT."""
    result=[];pos=0
    for match in re.finditer(r'([_^])\{([^}]+)\}',s):
        if match.start()>pos:result.append((s[pos:match.start()],1,0,True))
        val=match.group(2)
        result.append((val,.76,-22000 if match.group(1)=='_' else 35000,False))
        pos=match.end()
    if pos<len(s):result.append((s[pos:],1,0,True))
    clean=[]
    for content,factor,baseline,it in result:
        if factor!=1:
            clean.append((content,factor,baseline,it));continue
        for part in re.split(r'([A-Za-z]{3,})',content):
            if part:clean.append((part,factor,baseline,False if re.fullmatch(r'[A-Za-z]{3,}',part) else it))
    return clean

@lru_cache(maxsize=256)
def measure_font(size,bold,italic):
    directory=Path(os.environ.get('PRISMQUANT_FONT_DIR',str(Path.home()/'.local/share/figure-fonts/times-new-roman')))
    fname='Timesbi.TTF' if bold and italic else 'Timesbd.TTF' if bold else 'Timesi.TTF' if italic else 'Times.TTF'
    candidates={f.name.lower():f for f in directory.glob('*')}
    font_file=candidates.get(fname.lower(),directory/fname)
    if not font_file.exists():
        raise FileNotFoundError('Set PRISMQUANT_FONT_DIR to a directory containing Times New Roman TTF files.')
    return ImageFont.truetype(str(font_file),round(size*20))

def native_math(x,y,w,h,s,size,bold,color,align):
    y-=size*.15
    rr=runs_math(s);items=[];i=0
    while i<len(rr):
        t,f,b,it=rr[i];parts=[(t,f,b,it)]
        if f!=1 and i+1<len(rr) and rr[i+1][1]!=1 and b*rr[i+1][2]<0:
            parts.append(rr[i+1]);i+=1
        widths=[measure_font(size*ff,bold,ii).getlength(tt)/20 for tt,ff,bb,ii in parts]
        items.append((parts,max(widths)+(4 if f!=1 else 0)))
        i+=1
    total=sum(v for _,v in items)
    xx=x+(w-total)/2 if align=='center' else x+w-total if align=='right' else x
    for parts,width in items:
        for t,f,b,it in parts:
            tx=xx+(4 if f!=1 else 0)
            # Native editable textboxes with explicit position and physical font size.
            # No automatic superscript shrinkage in the office renderer.
            shift=(size*.25 if b<0 else -size*.35 if b>0 else 0)+(.30*size*(1-f))
            text(tx,y+shift,width+1.5,h,t,size*f,bold,color,'left',False,it)
        xx+=width
    return None

def text(x,y,w,h,s,size=26,bold=False,color='ink',align='left',mathtext=False,italic=False,name=''):
    if mathtext:
        return native_math(x,y,w,h,s,size,bold,color,align)
    box=shapes.add_textbox(u(x),u(y),u(w),u(h));box.name=name or 'Text: '+s[:45]
    tf=box.text_frame;tf.clear();tf.word_wrap=False;tf.auto_size=MSO_AUTO_SIZE.NONE
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    p=tf.paragraphs[0];p.alignment={'left':PP_ALIGN.LEFT,'center':PP_ALIGN.CENTER,'right':PP_ALIGN.RIGHT}[align]
    p.space_before=Pt(0);p.space_after=Pt(0)
    rr=runs_math(s) if mathtext else [(s,1,0,italic)]
    for t,factor,baseline,it in rr:
        run=p.add_run();run.text=t;run.font.name=FONT;run.font.size=Pt(size*SCALE*factor)
        run.font.bold=bold;run.font.italic=it;run.font.color.rgb=rgb(color)
        pr=run._r.get_or_add_rPr();pr.set('baseline',str(baseline));pr.set('lang','en-US');pr.set('dirty','0')
        for tag in ['a:latin','a:ea','a:cs']:
            e=pr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}'+tag.split(':')[1])
            if e is None:e=OxmlElement(tag);pr.append(e)
            e.set('typeface',FONT)
    text_records.append(dict(text=s,x=x,y=y,w=w,h=h,size=size,math=mathtext))
    return box

def box(x,y,w,h,label,fill='weight',color='ink',mathtext=False,size=27,linecolor='border'):
    rect(x,y,w,h,fill,linecolor,.95,5)
    return text(x+3,y+1,w-6,h-2,label,size=size,color=color,align='center',mathtext=mathtext)

def rotation(x,y,w=64,h=46,label='R_{2}'):
    return box(x,y,w,h,label,'lightpurple','purple',True,28,'purple')

def quant(x,y,w=40,h=46):return box(x,y,w,h,'A_{4}','lightblue','blue',True,27,'blue')
def weight(x,y,w,h,index):
    rect(x,y,w,h,'weight','border',.95,5)
    # Stacked, editable upper/lower indices instead of cumulative inline width.
    total=62;left=x+(w-total)/2
    text(left,y+1,27,34,'W',29,italic=True)
    text(left+27,y-1,36,22,'fold',22)
    text(left+27,y+20,40,19,index,22)

def badge(x,y,w,s,fill='lightblue',color='blue',size=22):
    rect(x,y,w,26,fill,None,r=5);text(x,y,w,26,s,size=size,color=color,align='center')
def plus(cx,cy):
    ellipse(cx-16,cy-16,32,32,'white','ink',1.2)
    line([(cx-7,cy),(cx+7,cy)],'ink',1.05,False)
    line([(cx,cy-7),(cx,cy+7)],'ink',1.05,False)
def panel(x,y,w,h,ident,title=None):
    rect(x,y,w,h,'panel','border',.8,10,name='Panel '+ident)
    panels.append(dict(id=ident,rect=[x,y,w,h],label=[x+18,y+18]))
    if title:text(x+18,y+10,w-36,35,title,30,True)

# A: original left-to-right Transformer context, with both residual paths.
group('01 Transformer backbone')
y=76
text(24,57,69,35,'tokens',25)
line([(91,y),(109,y)])
rect(109,44,157,64,'weight','border',.95,6)
text(113,45,149,29,'Embedding',26,align='center')
text(113,81,149,23,'R_{1} folded',22,color='purple',align='center',mathtext=True)
for x,label in [(299,'RMSNorm'),(718,'RMSNorm'),(1135,'RMSNorm')]:box(x-3,53,115,46,label,'white',size=24)
box(442,53,172,46,'Attention (a)','panel',size=27)
box(861,53,172,46,'FFN (b)','panel',size=27)
plus(666,y);plus(1083,y)
rect(1278,44,212,64,'weight','border',.95,6)
text(1282,45,204,29,'Output head',26,align='center')
text(1282,81,204,23,'R_{1}^{T} folded',22,color='purple',align='center',mathtext=True)
text(1516,57,65,35,'logits',25)
for a,b in [(266,299),(408,442),(614,650),(682,718),(827,861),(1033,1067),(1099,1135),(1244,1278),(1490,1515)]:line([(a,y),(b,y)])
line([(282,y),(282,25),(666,25),(666,60)])
line([(700,y),(700,25),(1083,25),(1083,60)])
dot(282,y);dot(700,y)
# The global scope is stated in Deploy; leave a clean top margin here.
text(1150,24,85,23,'final',22,color='muted',align='center')

# B: the same three cards, now with uniform title anchors and compact content.
CARD_Y=132;CARD_H=216
for ident,x,w,phase,title in [
    ('construct',24,366,'Construct','Spectral alignment'),
    ('represent',404,396,'Represent','Exploit the affine offset'),
    ('deploy',814,762,'Deploy','One principle, three sites')]:
    group('02 '+phase)
    panel(x,CARD_Y,w,CARD_H,ident)
    text(x+18,139,w-36,25,phase,23,True,color='purple')
    text(x+18,165,w-36,32,title,28,True)
    if ident=='construct':
        # Conceptual scatter: sparse unsaturated points, clear principal axes.
        cx,cy=133,253
        pts=[(-70,7),(-62,1),(-55,10),(-46,-1),(-43,15),(-35,1),(-28,-6),(-25,10),(-18,0),(-10,-12),(-9,10),(-3,-1),(4,9),(9,-8),(18,4),(22,-13),(29,-5),(33,10),(41,-8),(50,-15),(57,-7),(63,-19),(69,-10),(76,-22)]
        for i,(dx,dy) in enumerate(pts):dot(cx+dx,cy-.34*dx+dy*.38,1.25+(i%3)*.3,'blue' if i%4==0 else 'border')
        line([(54,287),(216,218)],'blue',1.2)
        line([(117,222),(148,282)],'border',.85)
        text(174,198,35,25,'v_{1}',25,mathtext=True)
        text(166,265,35,25,'v_{2}',25,mathtext=True)
        text(235,217,147,34,'R = H_{g}DΠG',26,mathtext=True)
        text(235,257,147,34,'G = I − WY^{T}',26,mathtext=True)
        text(47,307,318,28,'Σ  →  V_{k}  →  R',28,mathtext=True,align='center')
    elif ident=='represent':
        text(422,199,360,25,'Aligned group · same scale and range',22,color='muted')
        base=292;offset=46;vals=[46,59,50,62,54]
        for i,val in enumerate(vals):rect(435+i*23,base-val,14,val,'blue',None,r=0)
        for i,val in enumerate(vals):
            v=val-offset
            if v:rect(653+i*23,base-v,14,v,'blue',None,r=0)
            else:line([(653+i*23,base),(667+i*23,base)],'blue',1.2,False)
        line([(428,base),(552,base)],'border',.6,False)
        line([(646,base),(770,base)],'border',.6,False)
        line([(427,base-offset),(554,base-offset)],'green',1.0,False,True)
        text(426,294,141,26,'affine offset',22,color='green',align='center')
        line([(578,260),(623,260)],'muted',1.0)
        text(572,227,58,28,'− z',24,mathtext=True,align='center')
        # Exact same 16-pixel range bracket; translation cannot reduce range.
        for xx,yy in [(556,base-max(vals)),(772,base-(max(vals)-offset))]:
            line([(xx+4,yy),(xx,yy),(xx,yy+16),(xx+4,yy+16)],'purple',.8,False)
        text(641,294,139,26,'INT4 variation',22,color='blue',align='center')
        # Both brackets have the exact same extent under the same vertical scale.
        text(420,320,365,23,'shared level + within-group variation',22,align='center')
    else:
        # Shared scope is drawn explicitly: one tag per site, no loop icons.
        for yy,lab in [(202,'R_{1}'),(249,'R_{2}'),(296,'R_{4}')]:rotation(832,yy+1,49,39,lab)
        for yy in [207,218,229]:
            rect(908,yy,96,5,'lightpurple','purple',.45,0)
        line([(896,207),(896,234)],'purple',.65,False)
        for xx in [909,940,971]:rect(xx,257,23,27,'white','purple',.75,3)
        line([(897,251),(983,251)],'purple',.65,False)
        for xx in [920,951,982]:line([(xx,251),(xx,257)],'purple',.65,False)
        line([(902,316),(936,316)],'ink',.9)
        badge(939,303,73,'online',size=22)
        line([(1012,316),(1026,316)],'ink',.9)
        for yy,title,subtitle in [(201,'Global residual basis','Shared across layers'),(248,'Head-wise values','Shared across heads within a layer'),(295,'Down-projection input','Per layer')]:
            text(1040,yy,520,27,title,25,True)
            text(1040,yy+25,520,24,subtitle,23,color='muted')

# C1: QK scoring and attention-weighted V are two distinct, correctly connected nodes.
group('03 Attention expansion')
panel(24,362,739,360,'attention','(a)  Attention')
text(44,404,690,29,'Calibrated value rotation; inverse folded into output weights',23,color='muted')
text(40,505,27,29,'X',27,mathtext=True)
quant(73,502,40,46);line([(65,525),(73,525)])
line([(113,525),(127,525)],arrow=False)
line([(127,459),(127,611)],arrow=False);dot(127,525)
for yy,idx in [(436,'Q'),(512,'K'),(588,'V')]:
    line([(127,yy+23),(147,yy+23)])
    weight(147,yy,74,46,idx)
    line([(221,yy+23),(241,yy+23)])
box(241,436,77,46,'RoPE','white',size=25)
box(241,512,77,46,'RoPE','white',size=25)
rotation(247,588,65,46)
text(239,563,84,24,'online',22,color='blue',align='center')
line([(318,535),(341,535)]);line([(312,611),(341,611)])
box(341,512,105,46,'K cache','cache',size=25,linecolor='cacheedge')
box(341,588,105,46,'V cache','cache',size=25,linecolor='cacheedge')
text(333,558,121,26,'channel-wise',22,color='muted',align='center')
text(333,634,121,26,'token-wise',22,color='muted',align='center')
# The score block takes Q and K on separate ports.
line([(318,459),(470,459)])
line([(446,535),(456,535),(456,522),(470,522)])
rect(470,436,164,111,'white','border',.95,6)
text(479,444,146,32,'QK^{T} / √d_{h}',27,mathtext=True,align='center')
text(479,480,146,29,'causal mask',25,align='center')
text(479,511,146,28,'softmax',25,align='center')
line([(552,547),(552,588)])
text(566,550,83,31,'A_{attn}',24,mathtext=True)
line([(446,611),(484,611)])
box(484,588,123,46,'A_{attn} V','white',mathtext=True,size=27)
line([(607,611),(627,611)])
quant(627,588,40,46)
line([(667,611),(681,611)])
weight(681,588,64,46,'O')
line([(745,611),(755,611)])
text(44,690,697,26,'R_{2}^{−1} and residual basis folded into output weight',24,mathtext=True)

# C2: only gate passes SiLU; Pi and D Pi folds precede a conjugated online R4.
group('04 FFN expansion')
panel(779,362,797,360,'ffn','(b)  FFN')
text(799,404,743,29,'Signed permutation folded into SwiGLU projections',23,color='muted')
text(795,521,24,29,'X',27,mathtext=True)
quant(827,514,40,46);line([(818,537),(827,537)])
line([(867,537),(880,537)],arrow=False);dot(880,537)
line([(880,469),(880,605)],arrow=False)
line([(880,469),(899,469)]);line([(880,605),(899,605)])
weight(899,446,105,46,'gate');weight(899,582,105,46,'up')
text(899,496,105,26,'Π folded',23,color='purple',align='center')
text(899,632,105,26,'DΠ folded',23,color='purple',align='center')
line([(1004,469),(1023,469)])
box(1023,446,75,46,'SiLU','white',size=27)
line([(1098,469),(1115,469),(1115,520)])
line([(1004,605),(1115,605),(1115,554)])
ellipse(1098,520,34,34,'white','ink',1.05)
dot(1115,537,3.0,'ink')
rect(1148,447,303,212,'lightblue',None,r=10,name='Online structured transform group')
# The flow arrow is above the online-region background.
line([(1132,537),(1164,537)])
badge(1164,459,70,'online',fill='white',size=22)
text(1380,458,50,30,'R_{4}',27,color='purple',mathtext=True,align='right')
rotation(1164,514,78,46,"G_{4}'")
line([(1242,537),(1261,537)])
box(1261,514,81,46,'H_{128}','had','ink',True,27,'blue')
line([(1342,537),(1361,537)])
quant(1361,514,65,46)
text(1163,581,273,33,"G_{4}' = I − W'Y'^{T}",27,mathtext=True,align='center')
text(1163,620,273,26,'Conjugated correction',23,color='blue',align='center')
line([(1426,537),(1472,537)])
weight(1472,514,80,46,'D');line([(1552,537),(1567,537)])
text(1029,660,522,25,'Two-stage Triton implementation',24,color='blue',align='center')
text(799,690,755,26,'Inverse rotation and R_{1} output basis folded into W_{D}',24,mathtext=True)

# D: one compact legend strip, no enclosing panel and no decorative logo.
group('05 Legend')
line([(24,740),(1576,740)],'border',.7,False)
rotation(27,754,34,30,'R');text(71,752,203,34,'aligned rotation',23)
box(294,754,34,30,'H','had',mathtext=True,size=26,linecolor='blue');text(338,752,211,34,'block Hadamard',23)
box(563,754,34,30,'W','weight',mathtext=True,size=26);text(608,752,175,34,'folded weight',23)
quant(796,754,41,30);text(848,752,233,34,'activation quantization',23)
badge(1105,756,73,'online',size=22);text(1189,752,113,34,'execution',23)
box(1320,754,65,30,'K / V','cache',size=23,linecolor='cacheedge');text(1396,752,150,34,'cache',23)

# Preserve scientific context and drawing conventions in editable speaker notes.
slide.notes_slide.notes_text_frame.text=(
    'PrismQuant method refinement. Conceptual schematic; scatter and bars are illustrative, not measured data.\n'
    'Column convention: R = H_g D Pi G, G = I - W Y^T. With Q = D Pi and G prime = Q G Q^T, the online R4 path is H_128 G prime after Q is folded into gate/up weights. Gate uses Pi only; up uses D Pi.\n'
    'R1 is global and shared across layers. R2 is calibrated from tokens and KV heads pooled within each layer; the same head-dimensional rotation is applied to every head in that layer. R2 inverse compensation is folded into W_O.\n'
    'K cache is channel-wise after RoPE; V cache is token-wise after R2. Full-precision residual-cache details are intentionally omitted.\n'
    'The representation panel has identical scales and max-minus-min ranges before/after subtracting the affine offset. Range reduction comes from alignment; a shared component is retained in the representation. The conceptual shared level c_j need not equal the stored minimum z_j.\n'
    'Operator boxes indicate mathematical order, not kernel-launch boundaries. E17 v3 computes a rank-k projection followed by transform/correction and packing. No speed claims are made.\n'
    'Implementation source: https://github.com/ForeverBlue816/nar-offline-validation ; checked files: nar/e14_w4a4kv4.py, nar/activation_experiments.py, nar/fold_signed_permutation.py, nar/e17_v3.py.\n'
    'Every diagram object and all text are native editable PowerPoint elements. Groups can be ungrouped. English and mathematical text use Times New Roman. Superscripts and subscripts are individually positioned editable text objects, preserving legibility across office renderers.'
)
prs.core_properties.title='PrismQuant: aligned rotation and affine INT4 representation'
prs.core_properties.subject='Editable method schematic, refined from the supplied layout'
prs.core_properties.author=''

def main():
    p=argparse.ArgumentParser();p.add_argument('--out',type=Path,default=Path(__file__).parent)
    p.add_argument('--render',action='store_true');p.add_argument('--soffice',default='soffice')
    a=p.parse_args();a.out.mkdir(parents=True,exist_ok=True);(a.out/'qa').mkdir(exist_ok=True)
    # Clear inherited PowerPoint theme effects, including default shadows.
    def clear_effects(items):
        for sh in items:
            if sh.shape_type==6:clear_effects(sh.shapes)
            for pr in sh._element.xpath('./p:spPr'):
                eff=OxmlElement('a:effectLst');pr.append(eff)
            for eff in sh._element.xpath('./p:style/a:effectRef'):eff.set('idx','0')
    clear_effects(slide.shapes)
    stem='PrismQuant_method_refined';ppt=a.out/(stem+'.pptx');prs.save(ppt)
    (a.out/'qa'/'layout_source.json').write_text(json.dumps({'panels':panels,'texts':text_records},indent=2))
    print(ppt)
    if a.render:
        work=a.out/'qa'/'ppt_render';work.mkdir(exist_ok=True)
        subprocess.run([a.soffice,'-env:UserInstallation=file://'+str((work/'profile').resolve()),'--headless','--convert-to','pdf','--outdir',str(work),str(ppt)],check=True)
        import pymupdf as fitz
        source=fitz.open(work/(stem+'.pdf'));doc=fitz.open()
        width=183/25.4*72;page=doc.new_page(width=width,height=width/2)
        page.show_pdf_page(page.rect,source,0)
        pdf=a.out/(stem+'.pdf');doc.save(pdf,garbage=4,deflate=True)
        page.get_pixmap(dpi=600,alpha=False).save(a.out/(stem+'.png'))
        page.get_pixmap(dpi=150,alpha=False).save(a.out/'qa'/'paper_width_150dpi.png')
        page.get_pixmap(matrix=fitz.Matrix(1600/width,1600/width),alpha=False).save(a.out/'qa'/'preview_1600.png')
        print(pdf)

if __name__=='__main__':main()
